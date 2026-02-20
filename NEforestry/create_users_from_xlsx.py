import sqlite3
import hashlib
from openpyxl import load_workbook

DB_FILE = "users.db"
EXCEL_FILE = "recipients_test2.xlsx"

# ---------------------------
# Helper functions
# ---------------------------

def hash_password(password):
    return hashlib.sha256(str(password).encode()).hexdigest()

def create_users_table():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE,
            password_hash TEXT
        )
    """)
    conn.commit()
    conn.close()

def add_user(email, password):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    try:
        c.execute(
            "INSERT INTO users (email, password_hash) VALUES (?, ?)",
            (email, hash_password(password))
        )
        conn.commit()
        print(f"Added: {email}")
    except sqlite3.IntegrityError:
        print(f"Already exists: {email}")
    finally:
        conn.close()

# ---------------------------
# Main script
# ---------------------------

if __name__ == "__main__":

    create_users_table()

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Oletetaan että headerit ovat:
    # email | name | username | password
    for row in ws.iter_rows(min_row=2, values_only=True):
        email = row[0]
        password = row[3]

        if email and password:
            add_user(email.strip(), str(password).strip())

    print("All users processed.")